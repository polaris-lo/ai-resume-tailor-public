"""
ATS 友好格式简历生成器

将 LLM 返回的结构化 profile JSON 生成为 ATS 标准格式 .docx。

格式原则：
- 不使用表格、分栏、文本框
- 节标题使用 Heading 1 样式 + 下划线分隔
- 字体：宋体（中文）/ Times New Roman（英文数字）
- 正文 11pt，姓名 16pt
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT_CN = "宋体"
FONT_EN = "Times New Roman"
COLOR_BLACK = RGBColor(0, 0, 0)


class AtsGenerator:
    def __init__(self, profile: dict):
        self.profile = profile
        self.doc = Document()
        self._setup_document()

    def _setup_document(self):
        for section in self.doc.sections:
            section.top_margin    = Cm(2)
            section.bottom_margin = Cm(2)
            section.left_margin   = Cm(2.5)
            section.right_margin  = Cm(2.5)

        normal = self.doc.styles["Normal"]
        normal.font.name  = FONT_EN
        normal.font.size  = Pt(11)
        normal.font.color.rgb = COLOR_BLACK
        normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)

    # ── 辅助方法 ──────────────────────────────────────────────────────────────

    def _run_font(self, run, size_pt: float = 11, bold: bool = False):
        run.font.name  = FONT_EN
        run.font.size  = Pt(size_pt)
        run.font.bold  = bold
        run.font.color.rgb = COLOR_BLACK
        run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)

    def _add_section_heading(self, title: str):
        heading = self.doc.add_heading(title, level=1)
        heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = heading.runs[0] if heading.runs else heading.add_run(title)
        self._run_font(run, size_pt=13, bold=True)

        # 下划线分隔线
        p = self.doc.add_paragraph()
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), "000000")
        pBdr.append(bottom)
        pPr.append(pBdr)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after  = Pt(2)

    def _add_para(self, text: str, bold: bool = False, indent: bool = False):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after  = Pt(1)
        if indent:
            p.paragraph_format.left_indent = Cm(0.5)
        run = p.add_run(text)
        self._run_font(run, bold=bold)

    def _add_kv_para(self, key: str, value: str):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after  = Pt(1)
        r_key = p.add_run(f"{key}：")
        self._run_font(r_key, bold=True)
        r_val = p.add_run(value)
        self._run_font(r_val)

    def _add_bullet(self, text: str):
        self._add_para(f"• {text}", indent=True)

    # ── 各节内容 ──────────────────────────────────────────────────────────────

    def _write_basic_info(self):
        info = self.profile.get("basic_info", {})
        self._add_section_heading("基本信息")

        if info.get("name"):
            p = self.doc.add_paragraph()
            run = p.add_run(info["name"])
            self._run_font(run, size_pt=16, bold=True)

        for key, field in [("手机", "phone"), ("邮箱", "email")]:
            if info.get(field):
                self._add_kv_para(key, info[field])

        for item in info.get("other", []):
            if item.strip():
                self._add_para(item)

    def _write_education(self):
        edu_list = self.profile.get("education", [])
        if not edu_list:
            return
        self._add_section_heading("教育背景")

        for edu in edu_list:
            start  = edu.get("start_date", "")
            end    = edu.get("end_date", "")
            school = edu.get("school", "")
            degree = edu.get("degree", "")
            major  = edu.get("major", "")
            college = edu.get("college", "")

            date_range = f"{start}-{end}" if start or end else ""
            college_major = f"{college} {major}专业".strip() if college else f"{major}专业"
            parts = [p for p in [date_range, school, college_major, degree] if p]
            self._add_para("    ".join(parts), bold=True)

            for extra in edu.get("extras", []):
                if extra.strip():
                    self._add_para(f"    {extra}")

    def _write_skills(self):
        skills = self.profile.get("skills", {})
        if not skills:
            return
        self._add_section_heading("专业技能")
        for category, description in skills.items():
            self._add_kv_para(category, description)

    def _write_work_experience(self):
        jobs = self.profile.get("work_experience", [])
        if not jobs:
            return
        self._add_section_heading("实习经历")

        for job in jobs:
            start   = job.get("start_date", "")
            end     = job.get("end_date", "")
            company = job.get("company", "")
            title   = job.get("title", "")
            city    = job.get("city", "")

            date_range = f"{start}-{end}" if start or end else ""
            city_str   = f"（{city}）" if city else ""
            parts = [p for p in [date_range, f"{company}{city_str}", title] if p]
            self._add_para("    ".join(parts), bold=True)

            for bullet in job.get("bullets", []):
                if bullet.strip():
                    self._add_bullet(bullet)

            self.doc.add_paragraph()

    def _write_projects(self):
        projects = self.profile.get("projects", [])
        if not projects:
            return
        self._add_section_heading("项目经历")

        for proj in projects:
            start = proj.get("start_date", "")
            end   = proj.get("end_date", "")
            name  = proj.get("name", "")
            role  = proj.get("role", "")

            date_range = f"{start}-{end}" if start or end else ""
            parts = [p for p in [date_range, name, role] if p]
            self._add_para("    ".join(parts), bold=True)

            for bullet in proj.get("bullets", []):
                if bullet.strip():
                    self._add_bullet(bullet)

            self.doc.add_paragraph()

    def _write_awards(self):
        awards = self.profile.get("awards", [])
        if not awards:
            return
        self._add_section_heading("获奖情况")

        for award in awards:
            date = award.get("date", "")
            name = award.get("name", "")
            text = f"{date}    {name}" if date else name
            self._add_para(text)

    # ── 生成入口 ──────────────────────────────────────────────────────────────

    def generate(self, output_path: str | Path) -> Path:
        output_path = Path(output_path)
        self._write_basic_info()
        self._write_education()
        self._write_skills()
        self._write_work_experience()
        self._write_projects()
        self._write_awards()
        self.doc.save(str(output_path))
        return output_path
